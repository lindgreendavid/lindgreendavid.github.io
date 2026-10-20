"""Builds the downloadable Excel model and CSV from src/data/pharma/dataset.json.
All calculations in the workbook are live formulas (inputs in blue). A Checks sheet tests identities.
Run: python tools/pharma/build-workbook.py"""
import csv
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
ds = json.loads((ROOT / "src/data/pharma/dataset.json").read_text(encoding="utf-8"))
val = json.loads((ROOT / "src/data/pharma/validation.json").read_text(encoding="utf-8"))
ref = json.loads((ROOT / "src/data/pharma/model-reference.json").read_text(encoding="utf-8"))
OUT = ROOT / "public/downloads/pharma-valuation"
OUT.mkdir(parents=True, exist_ok=True)

BLUE = Font(color="0000FF")
BOLD = Font(bold=True)
HEAD = PatternFill("solid", fgColor="E1EEE6")
INPUT = PatternFill("solid", fgColor="FFF8DC")
YEARS = [2023, 2024, 2025]
NAMES = {k: v["name"] for k, v in ds["companies"].items()}

wb = Workbook()

# ---------------------------------------------------------------- README
ws = wb.active
ws.title = "README"
lines = [
    ("Pharma Valuation Explorer: model workbook", BOLD),
    ("Generated " + ds["generated"] + " from public SEC XBRL data (Forms 20-F) and ECB reference rates.", None),
    ("", None),
    ("What this is", BOLD),
    ("A transparent comparison of Novartis, AstraZeneca and GSK for FY2023-FY2025, and a DCF for AstraZeneca. It is an educational model, not investment advice or a price target.", None),
    ("", None),
    ("Sheets", BOLD),
    ("Dataset: every figure with its XBRL tag, SEC accession number and filing date. FX: ECB reference rates. Comparables: formulas over Dataset. DCF: inputs in blue, formulas elsewhere. Sensitivity: value per share across WACC and terminal growth. Checks: identity tests.", None),
    ("", None),
    ("Key definitions", BOLD),
    ("Revenue: AstraZeneca total revenue; GSK turnover; Novartis net sales plus other revenues. Profit: continuing operations (Novartis 2023 excludes the Sandoz distribution gain).", None),
    ("Debt: financial debt including lease liabilities. Net debt: debt less cash and, for Novartis, marketable securities and time deposits. Free cash flow: operating cash flow less purchases of PP&E (and, in the stricter version, intangible assets).", None),
    ("GSK is reported in GBP; USD translation uses ECB rates (annual average for flows, year-end for balances).", None),
    ("", None),
    ("Limits you should know", BOLD),
    ("The DCF bridge is partial: it omits contingent consideration, litigation and other provisions, deferred tax, other investments and any dividends declared. Terminal value is about three quarters of enterprise value in the base case. Scenario assumptions are illustrative, not forecasts.", None),
    ("Market capitalisation is not in the dataset: share-price data could not be obtained from a source whose terms allow redistribution. Enter your own in the web tool.", None),
    ("The reconciliation to filing pages is a consistency check (" + str(val["totals"]["reconciledFigures"]) + " figures), not an audit.", None),
]
for n, (text, font) in enumerate(lines, start=1):
    c = ws.cell(row=n, column=1, value=text)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if font:
        c.font = font
ws.column_dimensions["A"].width = 130

# ---------------------------------------------------------------- Dataset
wd = wb.create_sheet("Dataset")
head = ["Company", "Year", "Metric", "Value (millions, reporting currency)", "Currency", "Source tag", "SEC accession", "Filing date", "Filing index"]
wd.append(head)
rows_csv = [head]
for cid, co in ds["companies"].items():
    for y in YEARS:
        for metric, cell in co["years"][str(y)].items():
            accn = str(cell.get("accn", ""))
            first = accn.split(",")[0].strip()
            index = f"https://www.sec.gov/Archives/edgar/data/{int(co['cik'])}/{first.replace('-', '')}/" if first else ""
            row = [co["name"], y, metric, cell["value"], co["currency"], cell.get("tag", ""), accn, cell.get("filed", ""), index]
            wd.append(row)
            rows_csv.append(row)
for c in wd[1]:
    c.font = BOLD
    c.fill = HEAD
for col, w in zip("ABCDEFGHI", [14, 7, 28, 22, 10, 62, 28, 12, 70]):
    wd.column_dimensions[col].width = w
wd.freeze_panes = "A2"
(OUT / "pharma-dataset.csv").write_text("\n".join(",".join('"' + str(x).replace('"', '""') + '"' for x in r) for r in rows_csv), encoding="utf-8")

# ---------------------------------------------------------------- FX
wf = wb.create_sheet("FX")
wf.append(["Year", "USD per EUR (average)", "USD per EUR (year-end)", "GBP per EUR (average)", "GBP per EUR (year-end)", "USD per GBP (average)", "USD per GBP (year-end)", "Source"])
for i, y in enumerate(YEARS, start=2):
    wf.append([y, ds["fx"]["usdPerEurAvg"][str(y)], ds["fx"]["usdPerEurEnd"][str(y)], ds["fx"]["gbpPerEurAvg"][str(y)], ds["fx"]["gbpPerEurEnd"][str(y)], f"=B{i}/D{i}", f"=C{i}/E{i}", ds["fx"]["source"] if i == 2 else ""])
for c in wf[1]:
    c.font = BOLD
    c.fill = HEAD
for col in "ABCDEFG":
    wf.column_dimensions[col].width = 24
wf.column_dimensions["H"].width = 90

# ---------------------------------------------------------------- Comparables
wc = wb.create_sheet("Comparables")
cols = ["Company", "Year", "Currency", "Revenue", "Operating profit", "Operating margin", "Profit (continuing)", "Profit margin", "R&D", "R&D % revenue", "Operating cash flow", "Purchases of PP&E", "Purchases of intangibles", "FCF (CFO - PP&E)", "FCF after intangibles", "Cash", "Debt incl. leases", "Net debt", "Revenue (USD m)", "FCF after intangibles (USD m)"]
wc.append(cols)
r = 2
for cid, co in ds["companies"].items():
    for y in YEARS:
        sf = lambda m, r=r: f'SUMIFS(Dataset!$D:$D,Dataset!$A:$A,$A{r},Dataset!$B:$B,$B{r},Dataset!$C:$C,"{m}")'
        wc.append(
            [
                co["name"], y, co["currency"],
                f"={sf('revenue')}", f"={sf('operating_profit')}", f"=E{r}/D{r}", f"={sf('profit_continuing')}", f"=G{r}/D{r}",
                f"={sf('rnd')}", f"=I{r}/D{r}", f"={sf('cfo')}", f"={sf('capex_ppe')}", f"={sf('capex_intangibles')}",
                f"=K{r}-L{r}", f"=K{r}-L{r}-M{r}", f"={sf('cash')}",
                f'=IF({sf("borrowings_incl_leases")}<>0,{sf("borrowings_incl_leases")},{sf("debt_financial")}+{sf("lease_liabilities")})',
                f'=Q{r}-P{r}-{sf("liquid_investments")}',
                f'=D{r}*IF(C{r}="GBP",INDEX(FX!$F$2:$F$4,MATCH(B{r},FX!$A$2:$A$4,0)),1)',
                f'=O{r}*IF(C{r}="GBP",INDEX(FX!$F$2:$F$4,MATCH(B{r},FX!$A$2:$A$4,0)),1)',
            ]
        )
        for col in (6, 8, 10):
            wc.cell(row=r, column=col).number_format = "0.0%"
        for col in (4, 5, 7, 9, 11, 12, 13, 14, 15, 16, 17, 18, 19, 20):
            wc.cell(row=r, column=col).number_format = "#,##0"
        r += 1
for c in wc[1]:
    c.font = BOLD
    c.fill = HEAD
    c.alignment = Alignment(wrap_text=True, vertical="top")
for i in range(1, len(cols) + 1):
    wc.column_dimensions[get_column_letter(i)].width = 15
wc.freeze_panes = "D2"

# ---------------------------------------------------------------- DCF
wd2 = wb.create_sheet("DCF")
az = "AstraZeneca"
q = lambda m, y=2025: f'SUMIFS(Dataset!$D:$D,Dataset!$A:$A,"{az}",Dataset!$B:$B,{y},Dataset!$C:$C,"{m}")'
inputs = ref["inputs"]
rowsI = [
    ("AstraZeneca DCF (USD millions, valuation date 31 Dec 2025)", None, None),
    ("Blue cells are inputs. Everything else is a formula. Scenario assumptions are illustrative, not forecasts.", None, None),
    ("", None, None),
    ("Base year (fiscal)", 2025, "in"),
    ("Base revenue", f"={q('revenue')}", None),
    ("Base operating margin", f"={q('operating_profit')}/B5", None),
    ("Tax rate on operating profit", inputs["taxRate"], "in"),
    ("D&A, % of revenue", inputs["daPct"], "in"),
    ("Capex (PP&E + intangibles), % of revenue", inputs["capexPct"], "in"),
    ("Investment in working capital, % of revenue", inputs["nwcPct"], "in"),
    ("Operating margin in final forecast year", inputs["targetMargin"], "in"),
    ("WACC", inputs["wacc"], "in"),
    ("Terminal growth", inputs["terminalGrowth"], "in"),
]
for n, (a, b, kind) in enumerate(rowsI, start=1):
    wd2.cell(row=n, column=1, value=a)
    if b is not None:
        c = wd2.cell(row=n, column=2, value=b)
        if kind == "in":
            c.font = BLUE
            c.fill = INPUT
wd2["A1"].font = BOLD
for rr in (6, 7, 8, 9, 10, 11, 12, 13):
    wd2.cell(row=rr, column=2).number_format = "0.0%"
wd2["B5"].number_format = "#,##0"

wd2.cell(row=16, column=1, value="Forecast year (t)")
wd2.cell(row=17, column=1, value="Fiscal year")
wd2.cell(row=18, column=1, value="Revenue growth")
labels = ["Revenue", "Operating margin", "Operating profit (EBIT)", "NOPAT (EBIT x (1 - tax))", "Plus: D&A", "Less: capex (PP&E + intangibles)", "Less: investment in working capital", "Free cash flow to the firm", "Discount factor", "Present value of FCFF"]
for n, lab in enumerate(labels, start=19):
    wd2.cell(row=n, column=1, value=lab)
for t in range(1, 6):
    col = get_column_letter(1 + t)
    prev = get_column_letter(t)
    wd2[f"{col}16"] = t
    wd2[f"{col}17"] = f"=$B$4+{col}16"
    g = wd2[f"{col}18"]
    g.value = inputs["growth"][t - 1]
    g.font = BLUE
    g.fill = INPUT
    g.number_format = "0.0%"
    wd2[f"{col}19"] = f"=B5*(1+{col}18)" if t == 1 else f"={prev}19*(1+{col}18)"
    wd2[f"{col}20"] = f"=$B$6+($B$11-$B$6)*{col}16/5"
    wd2[f"{col}21"] = f"={col}19*{col}20"
    wd2[f"{col}22"] = f"={col}21*(1-$B$7)"
    wd2[f"{col}23"] = f"={col}19*$B$8"
    wd2[f"{col}24"] = f"={col}19*$B$9"
    wd2[f"{col}25"] = f"={col}19*$B$10"
    wd2[f"{col}26"] = f"={col}22+{col}23-{col}24-{col}25"
    wd2[f"{col}27"] = f"=1/(1+$B$12)^{col}16"
    wd2[f"{col}28"] = f"={col}26*{col}27"
    for rr in (19, 21, 22, 23, 24, 25, 26, 28):
        wd2[f"{col}{rr}"].number_format = "#,##0"
    wd2[f"{col}20"].number_format = "0.0%"
    wd2[f"{col}27"].number_format = "0.0000"
tail = [
    (30, "Sum of present values of FCFF", "=SUM(B28:F28)"),
    (31, "Terminal value at end of year 5 (Gordon growth)", "=F26*(1+B13)/(B12-B13)"),
    (32, "Present value of terminal value", "=B31*F27"),
    (33, "Enterprise value", "=B30+B32"),
    (34, "Terminal value as share of enterprise value", "=B32/B33"),
    (36, "Enterprise-to-equity bridge (31 Dec 2025 balance sheet)", None),
    (37, "Debt including lease liabilities", f"={q('borrowings_incl_leases')}"),
    (38, "Cash and cash equivalents", f"={q('cash')}"),
    (39, "Equity-accounted investments (book value)", f"={q('equity_method_investments')}"),
    (40, "Pension deficit (retirement benefit obligations, pre-tax)", f"={q('pension_liability')}"),
    (41, "Non-controlling interests (book value)", f"={q('nci_book')}"),
    (42, "Equity value", "=B33-B37+B38+B39-B40-B41"),
    (43, "Ordinary shares outstanding (millions)", f"={q('shares_outstanding_m')}"),
    (44, "Value per ordinary share (USD)", "=B42/B43"),
    (46, "Not included in the bridge: contingent consideration, litigation and other provisions, deferred tax, other investments, dividends declared. The bridge is partial.", None),
]
for rr, a, b in tail:
    wd2.cell(row=rr, column=1, value=a)
    if b:
        wd2.cell(row=rr, column=2, value=b)
        wd2.cell(row=rr, column=2).number_format = "#,##0.0" if rr in (43, 44) else ("0.0%" if rr == 34 else "#,##0")
wd2["A36"].font = BOLD
wd2["A33"].font = BOLD
wd2["A44"].font = BOLD
wd2.column_dimensions["A"].width = 62
for col in "BCDEF":
    wd2.column_dimensions[col].width = 14

# ---------------------------------------------------------------- Sensitivity
wsn = wb.create_sheet("Sensitivity")
wsn["A1"] = "Value per ordinary share (USD) across WACC (rows) and terminal growth (columns). Same FCFF as the DCF sheet."
wsn["A1"].font = BOLD
waccs = [0.06, 0.065, 0.07, 0.075, 0.08, 0.085, 0.09]
growths = [0.01, 0.015, 0.02, 0.025, 0.03]
wsn["A4"] = "WACC \\ g"
for j, g in enumerate(growths):
    c = wsn.cell(row=4, column=2 + j, value=g)
    c.number_format = "0.0%"
    c.font = BOLD
for i, w in enumerate(waccs):
    rr = 5 + i
    c = wsn.cell(row=rr, column=1, value=w)
    c.number_format = "0.0%"
    c.font = BOLD
    for j, g in enumerate(growths):
        col = get_column_letter(2 + j)
        f = (
            f"=IF($A{rr}<={col}$4,NA(),(SUMPRODUCT(DCF!$B$26:$F$26/(1+$A{rr})^DCF!$B$16:$F$16)"
            f"+DCF!$F$26*(1+{col}$4)/($A{rr}-{col}$4)/(1+$A{rr})^5"
            f"-DCF!$B$37+DCF!$B$38+DCF!$B$39-DCF!$B$40-DCF!$B$41)/DCF!$B$43)"
        )
        wsn.cell(row=rr, column=2 + j, value=f).number_format = "0.0"
wsn.column_dimensions["A"].width = 14

# ---------------------------------------------------------------- Checks
wk = wb.create_sheet("Checks")
wk.append(["Check", "Result", "Detail"])
checks = [
    ("Bridge identity: equity = EV - debt + cash + investments - pension - NCI", "=ABS(DCF!B42-(DCF!B33-DCF!B37+DCF!B38+DCF!B39-DCF!B40-DCF!B41))<0.001", "DCF!B42"),
    ("WACC exceeds terminal growth", "=DCF!B12>DCF!B13", "DCF!B12 vs DCF!B13"),
    ("Discount factors fall each year", "=AND(DCF!C27<DCF!B27,DCF!D27<DCF!C27,DCF!E27<DCF!D27,DCF!F27<DCF!E27)", "DCF!B27:F27"),
    ("Terminal value share of EV is between 0 and 100%", "=AND(DCF!B34>0,DCF!B34<1)", "DCF!B34"),
    ("Sensitivity centre (7.5% WACC, 2.0% growth) equals DCF value per share", "=ABS(Sensitivity!D8-DCF!B44)<0.001", "Sensitivity!D8"),
    ("Value per share matches the web model's reference value for the same inputs (within 0.01%)", f"=ABS(DCF!B44-{ref['valuePerShare']})/{ref['valuePerShare']}<0.0001", f"reference {ref['valuePerShare']:.4f} from tests/pharma"),
    ("Enterprise value matches the web model's reference value (within 0.01%)", f"=ABS(DCF!B33-{ref['enterpriseValue']})/{ref['enterpriseValue']}<0.0001", f"reference {ref['enterpriseValue']:.2f}"),
    ("Novartis 2025 net debt incl. leases less reported net debt (21,947) equals lease liabilities", '=ABS((Comparables!R4-21947)-SUMIFS(Dataset!$D:$D,Dataset!$A:$A,"Novartis",Dataset!$B:$B,2025,Dataset!$C:$C,"lease_liabilities"))<0.5', "Comparables row 4 (Novartis 2025)"),
]
for a, b, c in checks:
    wk.append([a, b, c])
wk.column_dimensions["A"].width = 95
wk.column_dimensions["B"].width = 12
wk.column_dimensions["C"].width = 45
for c in wk[1]:
    c.font = BOLD
    c.fill = HEAD

wb.save(OUT / "pharma-valuation-model.xlsx")
(OUT / "pharma-dataset.json").write_text((ROOT / "src/data/pharma/dataset.json").read_text(encoding="utf-8"), encoding="utf-8")
print("workbook, csv and json written to", OUT)
